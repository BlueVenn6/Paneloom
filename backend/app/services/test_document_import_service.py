import io

from docx import Document
from openpyxl import Workbook

from backend.app.services.document_import_service import DocumentImportError, import_document


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("第一幕：主角在雨夜收到神秘来信。")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "剧本节拍"
    sheet.append(["场景", "画面"])
    sheet.append(["P001", "主角推开门"])
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def run_smoke_tests() -> None:
    samples = [
        ("story.txt", "第一幕：主角醒来。\n第二幕：城市停电。".encode("utf-8"), "主角醒来"),
        ("story_gbk.txt", "中文 GBK 文本：角色进入旧楼。".encode("gb18030"), "角色进入旧楼"),
        ("story.html", "<h1>标题</h1><p>旁白进入。</p>".encode("utf-8"), "旁白进入"),
        ("story.rtf", r"{\rtf1\ansi 第一幕\par 主角看见信件}".encode("utf-8"), "主角看见信件"),
        ("story.docx", make_docx(), "神秘来信"),
        ("story.xlsx", make_xlsx(), "主角推开门"),
    ]
    for filename, content, expected in samples:
        result = import_document(filename, content)
        assert result["character_count"] > 0, filename
        assert result["text"].strip(), filename
        assert expected in result["text"], filename

    try:
        import_document("story.doc", b"\xd0\xcf\x11\xe0binary-word")
    except DocumentImportError:
        pass
    else:
        raise AssertionError("Unsupported .doc should fail clearly.")


if __name__ == "__main__":
    run_smoke_tests()
    print("document import smoke tests passed")
