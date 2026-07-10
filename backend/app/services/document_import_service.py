from __future__ import annotations

import csv
import io
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

MAX_IMPORT_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".srt", ".csv"}
SUPPORTED_EXTENSIONS = {
    *TEXT_EXTENSIONS,
    ".pdf",
    ".docx",
    ".epub",
    ".rtf",
    ".html",
    ".htm",
    ".xlsx",
}


class DocumentImportError(ValueError):
    pass


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"p", "br", "div", "section", "article", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")


def import_document(filename: str, content: bytes) -> dict:
    if not content:
        raise DocumentImportError("文件是空的，无法导入。")
    if len(content) > MAX_IMPORT_BYTES:
        raise DocumentImportError("文件超过 20MB，请先拆分或导出为较小的文本文件。")

    extension = Path(filename or "").suffix.lower()
    parser = _parser_for(extension, content)
    if not parser:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentImportError(f"暂不支持 {extension or '未知'} 格式。可导入格式：{supported}")

    text, warnings = parser(content)
    text = _clean_text(text)
    if not text:
        if extension == ".pdf":
            raise DocumentImportError("没有从 PDF 中提取到可选择文字。它可能是扫描版，请先 OCR 或导出为文本后再导入。")
        raise DocumentImportError("没有从文件中提取到有效文字，请检查文件内容或换一种导出格式。")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        text = text[:MAX_EXTRACTED_CHARACTERS]
        warnings.append("提取文本超过 500,000 字符，已安全截断；建议拆分为多个项目或章节。")

    return {
        "filename": filename,
        "extension": extension,
        "text": text,
        "character_count": len(text),
        "warnings": warnings,
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
    }


def _parser_for(extension: str, content: bytes) -> Callable[[bytes], tuple[str, list[str]]] | None:
    if extension in TEXT_EXTENSIONS:
        return _parse_plain_text
    if extension in {".html", ".htm"}:
        return _parse_html
    if extension == ".rtf":
        return _parse_rtf
    if extension == ".pdf" or content.startswith(b"%PDF"):
        return _parse_pdf
    if extension == ".docx":
        return _parse_docx
    if extension == ".epub":
        return _parse_epub
    if extension == ".xlsx":
        return _parse_xlsx
    return None


def _parse_plain_text(content: bytes) -> tuple[str, list[str]]:
    decoded, encoding = _decode_text(content)
    if _looks_like_csv(decoded):
        text = _csv_to_text(decoded)
    else:
        text = decoded
    warnings = [] if encoding in {"utf-8-sig", "utf-8"} else [f"已使用 {encoding} 编码读取。"]
    return text, warnings


def _parse_html(content: bytes) -> tuple[str, list[str]]:
    decoded, encoding = _decode_text(content)
    return _html_to_text(decoded), [] if encoding in {"utf-8-sig", "utf-8"} else [f"已使用 {encoding} 编码读取。"]


def _parse_rtf(content: bytes) -> tuple[str, list[str]]:
    decoded, encoding = _decode_text(content)
    text = _rtf_to_text(decoded)
    warnings = ["RTF 已按基础文本规则提取，复杂格式可能不会完全保留。"]
    if encoding not in {"utf-8-sig", "utf-8"}:
        warnings.append(f"已使用 {encoding} 编码读取。")
    return text, warnings


def _parse_pdf(content: bytes) -> tuple[str, list[str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentImportError("缺少 PDF 解析依赖 pypdf，请安装 backend/requirements.txt 后重试。") from exc

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise DocumentImportError("PDF 文件无法解析，可能已损坏或加密。") from exc

    if reader.is_encrypted:
        raise DocumentImportError("PDF 已加密，当前无法读取。请先解密或导出为文本。")

    pages: list[str] = []
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        if page_text.strip():
            pages.append(page_text)

    return "\n\n".join(pages), []


def _parse_docx(content: bytes) -> tuple[str, list[str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentImportError("缺少 Word 解析依赖 python-docx，请安装 backend/requirements.txt 后重试。") from exc

    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise DocumentImportError("DOCX 文件无法解析，可能已损坏或不是标准 Word 文档。") from exc

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), []


def _parse_epub(content: bytes) -> tuple[str, list[str]]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as exc:
        raise DocumentImportError("EPUB 文件无法解析，可能已损坏。") from exc

    if sum(item.file_size for item in archive.infolist()) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        archive.close()
        raise DocumentImportError("EPUB 解压后超过 100MB，请先拆分文件。")

    html_names = sorted(
        name
        for name in archive.namelist()
        if name.lower().endswith((".xhtml", ".html", ".htm"))
        and not name.lower().endswith(("nav.xhtml", "toc.xhtml"))
    )
    parts: list[str] = []
    for name in html_names:
        try:
            decoded, _ = _decode_text(archive.read(name))
        except Exception:
            continue
        text = _html_to_text(decoded)
        if text:
            parts.append(text)
        if sum(len(part) for part in parts) > MAX_EXTRACTED_CHARACTERS:
            break
    archive.close()
    return "\n\n".join(parts), []


def _parse_xlsx(content: bytes) -> tuple[str, list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentImportError("XLSX 文件无法解析，可能已损坏或不是标准 Excel 文档。") from exc

    parts: list[str] = []
    for worksheet in workbook.worksheets:
        parts.append(f"[{worksheet.title}]")
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                parts.append(" | ".join(values))
    workbook.close()
    return "\n".join(parts), []


def _decode_text(content: bytes) -> tuple[str, str]:
    for encoding in ["utf-8-sig", "utf-8"]:
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue

    candidates = ["gb18030", "big5", "shift_jis"]
    best_text = content.decode("latin-1", errors="replace")
    best_encoding = "latin-1"
    best_score = _text_score(best_text) - 30

    for encoding in candidates:
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        score = _text_score(text)
        if score > best_score:
            best_text = text
            best_encoding = encoding
            best_score = score

    return best_text, best_encoding


def _text_score(text: str) -> float:
    if not text:
        return 0
    replacement_penalty = text.count("\ufffd") * 8
    controls = sum(1 for char in text if ord(char) < 32 and char not in "\n\r\t")
    printable = sum(1 for char in text if char.isprintable() or char in "\n\r\t")
    cjk = sum(1 for char in text if "\u4e00" <= char <= "\u9fff")
    return printable + cjk * 0.4 - replacement_penalty - controls * 3


def _looks_like_csv(text: str) -> bool:
    first_lines = [line for line in text.splitlines()[:5] if line.strip()]
    return bool(first_lines) and sum("," in line or "\t" in line for line in first_lines) >= 2


def _csv_to_text(text: str) -> str:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return "\n".join(" | ".join(cell.strip() for cell in row if cell.strip()) for row in reader)


def _html_to_text(text: str) -> str:
    parser = _TextExtractor()
    parser.feed(text)
    return "\n".join(parser.parts)


def _rtf_to_text(text: str) -> str:
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\tab", "\t", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = text.replace("{", "").replace("}", "")
    return text


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
