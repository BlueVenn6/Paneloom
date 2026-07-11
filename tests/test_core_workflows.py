import io
import base64
import json
import tempfile
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfWriter
import requests
from fastapi import UploadFile

from api import openai_image_client
from backend.app import config
from backend.app.models.schemas import (
    ProjectCreate,
    ProjectUpdate,
    ProjectWorkspacePayload,
    ProjectActivityHeartbeat,
    ImageGenerationRequest,
    ReviewRequest,
    SchedulePlanRequest,
    SettingsPayload,
    ShotCreate,
    ShotUpdate,
    SocialFeedbackCreate,
    StoryOptimizationRequest,
)
from backend.app.routes import projects as project_routes
from backend.app.routes import settings as settings_routes
from backend.app.services import (
    export_service,
    bible_analysis_service,
    image_service,
    project_service,
    story_optimization_service,
)
from backend.app.services.document_import_service import DocumentImportError, import_document
from backend.app.services.schedule_service import (
    add_social_feedback,
    generate_review,
    generate_schedule,
    get_schedule,
    record_project_activity,
)
from backend.app.services.shotlist_service import (
    add_shot,
    apply_bible_to_shot_prompt,
    list_shots,
    select_shot_version,
    update_shot,
)


class IsolatedProjectTestCase(unittest.TestCase):
    def setUp(self):
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary_directory.name)
        self.projects_dir = self.root / "db" / "projects"
        self.projects_home = self.root / "projects-home"
        self.registry_path = self.root / "app-data" / "registry" / "projects.json"
        self.exports_dir = self.root / "exports"
        self.references_dir = self.root / "assets" / "references"
        self.generated_dir = self.root / "assets" / "generated"
        self.patchers = [
            patch.object(project_service, "PROJECTS_DIR", self.projects_dir),
            patch.object(project_service, "PROJECTS_HOME", self.projects_home),
            patch.object(project_service, "REGISTRY_PATH", self.registry_path),
        ]
        for patcher in self.patchers:
            patcher.start()

    def tearDown(self):
        for patcher in reversed(self.patchers):
            patcher.stop()
        self.temporary_directory.cleanup()

    def create_project(self, name="Test Project"):
        return project_service.create_project(ProjectCreate(name=name))

    def test_project_workspace_is_isolated_and_lifecycle_is_recoverable(self):
        first = self.create_project("First")
        second = self.create_project("Second")
        project_service.save_project_workspace(
            first["project_id"],
            ProjectWorkspacePayload(
                script="first project script",
                story_goal="stronger opening",
                selected_aspects=["square"],
            ),
        )

        self.assertEqual(
            project_service.get_project_workspace(first["project_id"])["script"],
            "first project script",
        )
        self.assertEqual(project_service.get_project_workspace(second["project_id"])["script"], "")

        renamed = project_service.update_project(
            first["project_id"],
            ProjectUpdate(name="Renamed", archived=True),
        )
        self.assertEqual(renamed["name"], "Renamed")
        self.assertTrue(renamed["archived"])

        duplicate = project_service.duplicate_project(first["project_id"])
        self.assertEqual(
            project_service.get_project_workspace(duplicate["project_id"])["script"],
            "first project script",
        )
        deleted = project_service.delete_project(duplicate["project_id"])
        self.assertTrue(deleted["deleted"])
        self.assertTrue(Path(deleted["recoverable_path"]).exists())
        self.assertIsNone(project_service.get_project(duplicate["project_id"]))

        with self.assertRaises(FileNotFoundError):
            project_service.get_project_dir(r"..\..\outside")

    def test_locked_shot_cannot_be_overwritten_edited_or_deleted(self):
        project = self.create_project("Lock Guard")
        shot = add_shot(project["project_id"], ShotCreate(prompt="first frame"))
        with self.assertRaisesRegex(ValueError, "saved image"):
            project_service.require_project(project["project_id"])
            from backend.app.services.shotlist_service import patch_shot

            patch_shot(project["project_id"], shot["shot_id"], ShotUpdate(locked=True))

        image_path = self.generated_dir / project["project_id"] / "frame.png"
        image_path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (16, 16), color=(20, 30, 40)).save(image_path)
        version = {
            "version_id": "v1",
            "output_path": str(image_path),
            "output_url": f"/assets/generated/{project['project_id']}/frame.png",
            "outputs": [],
        }
        update_shot(
            project["project_id"],
            shot["shot_id"],
            {
                "status": "done",
                "output_path": str(image_path),
                "output_url": version["output_url"],
                "versions": [version],
                "selected_version_id": "v1",
            },
        )
        from backend.app.services.shotlist_service import delete_shot, patch_shot

        locked = patch_shot(
            project["project_id"],
            shot["shot_id"],
            ShotUpdate(locked=True, approval_status="approved"),
        )
        self.assertTrue(locked["locked"])
        self.assertEqual(locked["approval_status"], "approved")
        with self.assertRaisesRegex(ValueError, "Unlock"):
            patch_shot(project["project_id"], shot["shot_id"], ShotUpdate(prompt="overwrite"))
        with self.assertRaisesRegex(ValueError, "Unlock"):
            select_shot_version(project["project_id"], shot["shot_id"], "v1")
        with self.assertRaisesRegex(ValueError, "Unlock"):
            delete_shot(project["project_id"], shot["shot_id"])

        blocked = image_service.generate_shot_image(project["project_id"], shot["shot_id"])
        self.assertEqual(blocked["error_code"], "shot_locked")
        batch = image_service.generate_project_batch(project["project_id"], only_pending=False)
        self.assertEqual(batch["processed"], 0)

        unlocked = patch_shot(
            project["project_id"],
            shot["shot_id"],
            ShotUpdate(locked=False, approval_status="draft"),
        )
        self.assertFalse(unlocked["locked"])

    def test_concurrent_shot_additions_do_not_overwrite_each_other(self):
        project = self.create_project("Concurrent")

        def create(index):
            return add_shot(
                project["project_id"],
                ShotCreate(scene_id=f"P{index:03d}", prompt=f"Panel {index}"),
            )

        with ThreadPoolExecutor(max_workers=8) as executor:
            list(executor.map(create, range(40)))

        shots = list_shots(project["project_id"])
        self.assertEqual(len(shots), 40)
        self.assertEqual(len({shot["shot_id"] for shot in shots}), 40)

    def test_low_capacity_schedule_returns_overload_instead_of_hanging(self):
        project = self.create_project("Schedule")
        result = generate_schedule(
            project["project_id"],
            SchedulePlanRequest(
                start_date="2026-07-10",
                target_publish_date="2026-07-17",
                daily_available_hours=0.25,
                total_content_items=1,
                panels_per_content=120,
            ),
        )
        self.assertTrue(result["tasks"])
        self.assertTrue(result["plan"]["summary"]["overloaded_days"])

    def test_automatic_progress_rejects_legacy_import_derivatives_as_current(self):
        project = self.create_project("Progress")
        project_id = project["project_id"]
        project_service.save_project_workspace(
            project_id,
            {
                "script": "A newly imported script",
                "script_import": {"filename": "new.pdf", "warnings": []},
            },
        )
        project_service.save_project_bible(
            project_id,
            {
                "characters": "Old character",
                "locations": "Old location",
                "world_notes": "Old world",
            },
        )
        add_shot(project_id, ShotCreate(scene_id="OLD-001", prompt="Old prompt"))

        stages = {item["stage"]: item for item in get_schedule(project_id)["project_progress"]}
        self.assertEqual(stages["script"]["status"], "done")
        self.assertEqual(stages["bible"]["status"], "doing")
        self.assertEqual(stages["storyboard"]["status"], "blocked")

    def test_project_activity_is_automatic_and_export_progress_never_exceeds_total(self):
        project = self.create_project("Activity")
        project_id = project["project_id"]
        first = record_project_activity(
            project_id,
            ProjectActivityHeartbeat(seconds=30, page="storyboard"),
        )
        second = record_project_activity(
            project_id,
            ProjectActivityHeartbeat(seconds=30, page="schedule"),
        )
        self.assertEqual(first["total_seconds"], 30)
        self.assertEqual(second["total_seconds"], 60)
        export_dir = project_service.get_project_export_dir(project_id)
        (export_dir / "one.json").write_text("{}", encoding="utf-8")
        (export_dir / "two.csv").write_text("a,b", encoding="utf-8")
        schedule = get_schedule(project_id)
        self.assertEqual(schedule["activity"]["total_seconds"], 60)
        export_stage = next(item for item in schedule["project_progress"] if item["stage"] == "export")
        self.assertEqual((export_stage["completed"], export_stage["total"]), (1, 1))
        review = generate_review(
            project_id,
            ReviewRequest(period="week", anchor_date="2026-07-11"),
        )
        self.assertEqual(review["automatic_active_minutes"], 1)
        self.assertEqual(review["progress_source"], "project_state")

    def test_native_desktop_save_writes_the_selected_file_atomically(self):
        project = self.create_project("Desktop Save")
        destination = self.root / "exports outside project" / "优化稿.txt"
        upload = UploadFile(filename="optimized.txt", file=io.BytesIO("真实内容".encode("utf-8")))
        with patch.object(project_routes, "_desktop_save_dialog", return_value=str(destination)):
            response = project_routes.save_desktop_export(
                project["project_id"],
                upload,
                "optimized.txt",
            )
        self.assertTrue(response.ok)
        self.assertEqual(destination.read_text(encoding="utf-8"), "真实内容")
        self.assertFalse(list(destination.parent.glob("*.tmp")))

    def test_ai_reference_plan_creates_separate_production_assets(self):
        project = self.create_project("Asset Plan")
        provider_text = json.dumps(
            {
                "assets": [
                    {
                        "asset_id": "hero-turnaround",
                        "category": "character_turnaround",
                        "name": "主角三视图",
                        "purpose": "固定发型服装和体型",
                        "image_prompt": "Production character turnaround sheet, neutral background, no story sequence, no text",
                    },
                    {
                        "asset_id": "memory-umbrella",
                        "category": "prop",
                        "name": "记忆雨伞",
                        "purpose": "固定关键道具",
                        "image_prompt": "Isolated production prop reference, orthographic views, no story sequence, no text",
                    },
                ]
            },
            ensure_ascii=False,
        )
        with (
            patch.object(image_service, "is_offline_mode", return_value=False),
            patch.object(image_service, "get_active_settings", return_value={}),
            patch.object(image_service, "validate_text_settings", return_value=[]),
            patch.object(image_service, "get_text_api_config", return_value={"provider": "test", "model": "test"}),
            patch.object(image_service, "generate_text", return_value={"ok": True, "text": provider_text}),
        ):
            result = image_service.plan_reference_assets(
                project["project_id"],
                script="全新雨伞故事",
                categories=["main_characters", "props"],
            )
        self.assertTrue(result["ok"])
        self.assertEqual(result["count"], 2)
        self.assertEqual({item["category"] for item in result["plan"]}, {"character_turnaround", "prop"})
        progress = []

        def fake_image(_settings, *, output_path, **_kwargs):
            Path(output_path).write_bytes(b"image")
            return {"ok": True, "status": "done", "output_path": str(output_path)}

        with (
            patch.object(image_service, "is_offline_mode", return_value=False),
            patch.object(image_service, "get_active_settings", return_value={}),
            patch.object(image_service, "validate_api_settings", return_value=[]),
            patch.object(image_service, "generate_image", side_effect=fake_image),
        ):
            generated = image_service.generate_reference_concepts(
                project["project_id"],
                plan=result["plan"],
                progress_callback=lambda **value: progress.append(value),
            )
        self.assertEqual((generated["planned"], generated["completed"], generated["failed"]), (2, 2, 0))
        self.assertEqual(progress[-1]["completed"], 2)
        listed = image_service.list_reference_images(project["project_id"])
        self.assertEqual({item["name"] for item in listed}, {"主角三视图", "记忆雨伞"})

    def test_new_script_bible_analysis_excludes_old_story_entities(self):
        project = self.create_project("Fresh Bible")
        project_id = project["project_id"]
        project_service.save_project_bible(
            project_id,
            {
                "visual_style": "User-selected ink style",
                "characters": "Old character that must not leak",
                "locations": "Old location that must not leak",
                "world_notes": "Old world that must not leak",
            },
        )
        provider_result = json.dumps(
            {
                "visual_style": "User-selected ink style",
                "characters": "New umbrella repairer",
                "locations": "Rainy workshop",
                "props": "Memory umbrella",
                "world_notes": "Umbrellas remember lies",
                "prompt_prefix": "Stable ink comic",
                "negative_prompt": "watermark",
            }
        )
        with (
            patch.object(bible_analysis_service, "is_offline_mode", return_value=False),
            patch.object(bible_analysis_service, "get_active_settings", return_value={}),
            patch.object(bible_analysis_service, "validate_text_settings", return_value=[]),
            patch.object(
                bible_analysis_service,
                "get_text_api_config",
                return_value={"provider": "test", "model": "test"},
            ),
            patch.object(
                bible_analysis_service,
                "generate_text",
                return_value={"ok": True, "text": provider_result},
            ) as provider,
        ):
            result = bible_analysis_service.analyze_project_bible(
                project_id,
                script="A brand-new umbrella story",
                use_ai=True,
                replace_existing=True,
            )

        self.assertTrue(result["ok"])
        self.assertNotIn("Old character", provider.call_args.kwargs["user_prompt"])
        self.assertNotIn("Old character", json.dumps(result["bible"]))
        self.assertEqual(result["bible"]["characters"], "New umbrella repairer")

    def test_review_uses_latest_social_snapshot_and_persists(self):
        project = self.create_project("Review")
        common = {
            "content_id": "EP001",
            "platform": "Platform A",
            "publish_date": "2026-07-08",
        }
        add_social_feedback(
            project["project_id"],
            SocialFeedbackCreate(**common, snapshot_date="2026-07-09", views=100),
        )
        add_social_feedback(
            project["project_id"],
            SocialFeedbackCreate(**common, snapshot_date="2026-07-10", views=250),
        )
        add_social_feedback(
            project["project_id"],
            SocialFeedbackCreate(
                content_id="EP002",
                platform="Platform B",
                publish_date="2026-07-08",
                snapshot_date="2026-07-10",
                views=400,
            ),
        )
        review = generate_review(
            project["project_id"],
            ReviewRequest(period="week", anchor_date="2026-07-10"),
        )
        self.assertEqual(review["views"], 650)
        self.assertEqual(review["posts_logged"], 2)
        self.assertEqual(review["best_platform"], "Platform B")
        self.assertEqual(get_schedule(project["project_id"])["reviews"]["week"]["views"], 650)

    def test_exports_are_real_and_parseable(self):
        project = self.create_project("Exports")
        shot = add_shot(
            project["project_id"],
            ShotCreate(
                scene_id="P001",
                duration=2.5,
                visual_description="A door opens",
                prompt="A cinematic panel of a door opening",
                revision_notes="Narration line",
            ),
        )
        generated = project_service.get_project_generated_dir(project["project_id"]) / "panel.png"
        Image.new("RGB", (16, 16), color=(20, 30, 40)).save(generated)
        relative = project_service.to_project_relative_path(project["project_id"], generated)
        update_shot(
            project["project_id"],
            shot["shot_id"],
            {
                "status": "done",
                "output_path": relative,
                "output_url": project_service.project_file_url(
                    project["project_id"], "generated", generated.name
                ),
            },
        )
        paths = {}
        for export_format in ["json", "csv", "xlsx", "srt", "html"]:
            result = export_service.export_project(project["project_id"], export_format)
            path = project_service.resolve_project_path(project["project_id"], result["path"])
            self.assertTrue(path.exists())
            self.assertGreater(path.stat().st_size, 0)
            paths[export_format] = path

        workbook = load_workbook(paths["xlsx"], read_only=True)
        self.assertEqual(workbook.active["B2"].value, "P001")
        workbook.close()
        self.assertIn("00:00:02,500", paths["srt"].read_text(encoding="utf-8"))
        html = paths["html"].read_text(encoding="utf-8")
        self.assertIn("A door opens", html)
        self.assertIn('src="../generated/panel.png"', html)
        self.assertEqual(
            json.loads(paths["json"].read_text(encoding="utf-8"))["shots"][0]["scene_id"],
            "P001",
        )

    def test_reference_upload_requires_a_real_supported_image_and_avoids_overwrite(self):
        project = self.create_project("References")
        buffer = io.BytesIO()
        Image.new("RGB", (32, 32), color=(20, 120, 200)).save(buffer, format="PNG")
        first = image_service.save_reference_image(
            project["project_id"], "hero.png", buffer.getvalue()
        )
        second = image_service.save_reference_image(
            project["project_id"], "hero.png", buffer.getvalue()
        )
        self.assertNotEqual(first["filename"], second["filename"])
        self.assertTrue(
            project_service.resolve_project_path(project["project_id"], first["path"]).exists()
        )
        with self.assertRaises(ValueError):
            image_service.save_reference_image(project["project_id"], "fake.png", b"not an image")

    def test_generated_version_can_be_selected(self):
        project = self.create_project("Versions")
        shot = add_shot(project["project_id"], ShotCreate(prompt="Panel"))
        first_path = project_service.get_project_dir(project["project_id"]) / "first.png"
        second_path = project_service.get_project_dir(project["project_id"]) / "second.png"
        first_path.write_bytes(b"first")
        second_path.write_bytes(b"second")
        versions = [
            {
                "version_id": "v1",
                "output_path": str(first_path),
                "output_url": "/assets/generated/first.png",
                "outputs": [{"output_url": "/assets/generated/first.png"}],
            },
            {
                "version_id": "v2",
                "output_path": str(second_path),
                "output_url": "/assets/generated/second.png",
                "outputs": [{"output_url": "/assets/generated/second.png"}],
            },
        ]
        update_shot(project["project_id"], shot["shot_id"], {"versions": versions})
        selected = select_shot_version(project["project_id"], shot["shot_id"], "v1")
        self.assertEqual(selected["selected_version_id"], "v1")
        self.assertEqual(selected["output_url"], "/assets/generated/first.png")

    def test_corrupt_shot_json_is_not_silently_overwritten(self):
        project = self.create_project("Corrupt")
        path = project_service.get_project_dir(project["project_id"]) / "shots.json"
        path.write_text("{broken", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "not overwritten"):
            list_shots(project["project_id"])
        self.assertEqual(path.read_text(encoding="utf-8"), "{broken")

    def test_partial_multi_aspect_generation_keeps_successful_output(self):
        project = self.create_project("Partial Generation")
        shot = add_shot(project["project_id"], ShotCreate(prompt="Panel"))
        calls = 0

        def fake_generate(settings, prompt, size, quality, output_path):
            nonlocal calls
            calls += 1
            if calls == 1:
                Image.new("RGB", (1024, 1024), color=(30, 40, 50)).save(output_path)
                return {
                    "ok": True,
                    "status": "done",
                    "output_path": str(output_path),
                    "output_url": "",
                    "download_status": "saved_b64_json",
                    "download_error": "",
                    "endpoint": "https://example.invalid/images/generations",
                    "model": "test-image",
                    "size": size,
                    "quality": quality or "low",
                    "status_code": 200,
                }
            return {
                "ok": False,
                "status": "failed",
                "error": "rate_limited",
                "error_code": "rate_limited",
                "error_message": "second aspect failed",
                "status_code": 429,
            }

        active_settings = {
            "provider_type": "custom",
            "base_url": "https://example.invalid/v1",
            "api_key": "test-key",
            "image_model": "test-image",
            "image_size": "1024x1024",
            "image_quality": "low",
            "image_result_mode": "auto",
        }
        with (
            patch.object(image_service, "get_active_settings", return_value=active_settings),
            patch.object(image_service, "generate_image", side_effect=fake_generate),
        ):
            result = image_service.generate_shot_image(
                project["project_id"],
                shot["shot_id"],
                ImageGenerationRequest(aspect_presets=["square", "portrait_4_5"]),
            )

        self.assertEqual(result["status"], "failed")
        self.assertEqual(len(result["outputs"]), 1)
        stored = list_shots(project["project_id"])[0]
        self.assertTrue(stored["output_url"])
        self.assertFalse(Path(stored["output_path"]).is_absolute())
        self.assertTrue(
            project_service.resolve_project_path(project["project_id"], stored["output_path"]).is_file()
        )
        self.assertEqual(stored["versions"][-1]["status"], "partial")

    def test_workspace_aspect_overrides_legacy_settings_image_size(self):
        project = self.create_project("Workspace Aspect")
        project_service.save_project_workspace(
            project["project_id"],
            {"selected_aspects": ["square"]},
        )
        shot = add_shot(project["project_id"], ShotCreate(prompt="Square panel"))
        observed_sizes = []

        def fake_generate(settings, prompt, size, quality, output_path):
            observed_sizes.append(size)
            Image.new("RGB", (1024, 1024), color=(30, 40, 50)).save(output_path)
            return {
                "ok": True,
                "status": "done",
                "output_path": str(output_path),
                "output_url": "",
                "download_status": "saved_b64_json",
                "download_error": "",
                "endpoint": "https://example.invalid/images/generations",
                "model": "test-image",
                "size": size,
                "quality": quality or "low",
                "status_code": 200,
            }

        legacy_settings = {
            "provider_type": "custom",
            "base_url": "https://example.invalid/v1",
            "api_key": "test-key",
            "image_model": "test-image",
            "image_size": "64x64",
            "image_quality": "low",
            "image_result_mode": "auto",
        }
        with (
            patch.object(image_service, "get_active_settings", return_value=legacy_settings),
            patch.object(image_service, "generate_image", side_effect=fake_generate),
        ):
            result = image_service.generate_shot_image(
                project["project_id"],
                shot["shot_id"],
                ImageGenerationRequest(),
            )

        self.assertEqual(result["status"], "done")
        self.assertEqual(observed_sizes, ["1024x1024"])
        self.assertEqual(result["outputs"][0]["size"], "1024x1024")
        self.assertEqual(result["outputs"][0]["final_size"], "1024x1024")

    def test_bible_and_selected_reference_are_used_by_generation(self):
        project = self.create_project("Continuity Generation")
        project_service.save_project_bible(
            project["project_id"],
            {
                "characters": "Hero wears a red coat",
                "locations": "Rainy station platform",
                "props": "red umbrella",
                "visual_style": "clean manga line art",
                "prompt_prefix": "continuity fixture",
            },
        )
        shot = add_shot(
            project["project_id"],
            ShotCreate(prompt="Hero waits for the train"),
        )
        buffer = io.BytesIO()
        Image.new("RGB", (32, 32), color=(120, 30, 40)).save(buffer, format="PNG")
        reference = image_service.save_reference_image(
            project["project_id"],
            "hero-reference.png",
            buffer.getvalue(),
        )
        bible_prompt, _ = apply_bible_to_shot_prompt(
            shot,
            project_service.get_project_bible(project["project_id"]),
        )
        self.assertIn("Hero wears a red coat", bible_prompt)
        self.assertIn("Rainy station platform", bible_prompt)
        self.assertIn("red umbrella", bible_prompt)
        self.assertIn("clean manga line art", bible_prompt)

        observed = {}

        def fake_edit(config, prompt, image_path, size, quality, output_path):
            observed["prompt"] = prompt
            observed["image_paths"] = [Path(item) for item in image_path]
            Image.new("RGB", (64, 64), color=(40, 60, 90)).save(output_path)
            return {
                "ok": True,
                "status": "done",
                "output_path": str(output_path),
                "output_url": "",
                "download_status": "saved_b64_json",
                "download_error": "",
                "endpoint": "https://example.invalid/images/edits",
                "model": "controlled-image",
                "size": size,
                "quality": quality,
                "status_code": 200,
            }

        settings = {
            "provider_type": "custom",
            "base_url": "https://example.invalid/v1",
            "api_key": "controlled-test-key",
            "image_model": "controlled-image",
            "image_size": "64x64",
            "image_quality": "low",
            "image_result_mode": "auto",
        }
        with (
            patch.object(image_service, "get_active_settings", return_value=settings),
            patch.object(image_service, "edit_image", side_effect=fake_edit) as edit_call,
        ):
            result = image_service.generate_shot_image(
                project["project_id"],
                shot["shot_id"],
                ImageGenerationRequest(
                    size="64x64",
                    quality="low",
                    reference_filenames=[reference["filename"]],
                ),
            )

        self.assertEqual(result["status"], "done")
        edit_call.assert_called_once()
        self.assertTrue(observed["image_paths"])
        self.assertTrue(all(path.is_file() for path in observed["image_paths"]))
        self.assertIn("Hero wears a red coat", observed["prompt"])
        self.assertIn("Rainy station platform", observed["prompt"])


class StoryOptimizationTests(unittest.TestCase):
    def _run_with_response(self, response_text: str) -> dict:
        with (
            patch.object(story_optimization_service, "get_active_settings", return_value={"text_model": "test"}),
            patch.object(story_optimization_service, "validate_text_settings", return_value=[]),
            patch.object(story_optimization_service, "is_offline_mode", return_value=False),
            patch.object(
                story_optimization_service,
                "get_text_api_config",
                return_value={"provider": "test", "model": "test-model"},
            ),
            patch.object(
                story_optimization_service,
                "generate_text",
                return_value={"ok": True, "text": response_text},
            ),
        ):
            return story_optimization_service.optimize_story(
                StoryOptimizationRequest(
                    script="Original story",
                    use_ai=True,
                    output_language="zh-CN",
                )
            )

    def test_real_ai_result_cannot_reuse_the_original_script(self):
        result = self._run_with_response(
            json.dumps({"logline": "Same", "improved_script": "Original story"})
        )
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "unchanged_model_output")

    def test_real_ai_result_has_provenance_and_no_english_template_fill(self):
        result = self._run_with_response(
            json.dumps(
                {
                    "logline": "Changed",
                    "improved_script": "A genuinely revised story",
                    "risks": [{"risk": "Visual drift", "solution": "Lock references"}],
                }
            )
        )
        self.assertTrue(result["ok"])
        self.assertEqual(result["source"], "ai")
        self.assertEqual(result["model"], "test-model")
        self.assertEqual(result["output_language"], "zh-CN")
        self.assertEqual(result["storyboard_strategy"], [])
        self.assertEqual(result["risks"], [{"risk": "Visual drift", "solution": "Lock references"}])
        self.assertTrue(result["input_sha256"])


class ImageClientBoundaryTests(unittest.TestCase):
    def test_request_construction_and_atomic_result_persistence(self):
        image = io.BytesIO()
        Image.new("RGB", (8, 8), color=(10, 20, 30)).save(image, format="PNG")

        class Response:
            status_code = 200
            text = ""

            @staticmethod
            def json():
                return {"data": [{"b64_json": base64.b64encode(image.getvalue()).decode("ascii")}]}

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "generated.png"
            with patch.object(openai_image_client.requests, "post", return_value=Response()) as post:
                result = openai_image_client.generate_image(
                    {
                        "api_key": "controlled-test-key",
                        "base_url": "http://127.0.0.1:9/v1",
                        "image_model": "controlled-model",
                        "image_size": "64x64",
                        "image_quality": "low",
                        "image_result_mode": "prefer_b64_json",
                    },
                    "one test panel",
                    output_path=output,
                )

            request = post.call_args
            self.assertEqual(request.kwargs["json"]["model"], "controlled-model")
            self.assertEqual(request.kwargs["json"]["prompt"], "one test panel")
            self.assertEqual(request.kwargs["json"]["size"], "64x64")
            self.assertEqual(request.kwargs["json"]["quality"], "low")
            self.assertEqual(request.kwargs["json"]["response_format"], "b64_json")
            self.assertEqual(request.kwargs["timeout"], (30, openai_image_client.IMAGE_REQUEST_TIMEOUT_SECONDS))
            self.assertTrue(result["ok"])
            self.assertTrue(output.is_file())
            with Image.open(output) as persisted:
                self.assertEqual(persisted.size, (8, 8))
            self.assertEqual(list(output.parent.glob(".*.tmp")), [])

    def test_timeout_is_classified_without_a_false_output(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "must-not-exist.png"
            with patch.object(
                openai_image_client.requests,
                "post",
                side_effect=requests.ReadTimeout("controlled timeout"),
            ):
                result = openai_image_client.generate_image(
                    {
                        "api_key": "controlled-test-key",
                        "base_url": "http://127.0.0.1:9/v1",
                        "image_model": "controlled-model",
                    },
                    "timeout test",
                    output_path=output,
                )
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "request_timeout")
        self.assertFalse(output.exists())


class DocumentImportTests(unittest.TestCase):
    def test_supported_document_matrix_and_scanned_pdf_error(self):
        document = Document()
        document.add_paragraph("DOCX scene")
        docx = io.BytesIO()
        document.save(docx)

        workbook = Workbook()
        workbook.active.append(["XLSX scene"])
        xlsx = io.BytesIO()
        workbook.save(xlsx)

        epub = io.BytesIO()
        with zipfile.ZipFile(epub, "w") as archive:
            archive.writestr("chapter.xhtml", "<html><body><p>EPUB scene</p></body></html>")

        samples = [
            ("story.txt", b"TXT scene", "TXT scene"),
            ("story.md", b"# MD scene", "MD scene"),
            ("story.srt", b"1\n00:00:00,000 --> 00:00:01,000\nSRT scene", "SRT scene"),
            ("story.csv", b"scene,action\nP001,CSV scene\nP002,next", "CSV scene"),
            ("story.html", b"<p>HTML scene</p>", "HTML scene"),
            ("story.rtf", b"{\\rtf1 RTF scene}", "RTF scene"),
            ("story.epub", epub.getvalue(), "EPUB scene"),
            ("story.docx", docx.getvalue(), "DOCX scene"),
            ("story.xlsx", xlsx.getvalue(), "XLSX scene"),
        ]
        for filename, content, expected in samples:
            with self.subTest(filename=filename):
                self.assertIn(expected, import_document(filename, content)["text"])

        pdf = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(pdf)
        with self.assertRaisesRegex(DocumentImportError, "OCR"):
            import_document("scan.pdf", pdf.getvalue())


class SettingsSecurityTests(unittest.TestCase):
    def test_both_keys_are_masked_and_failed_connection_is_not_success(self):
        with tempfile.TemporaryDirectory() as directory:
            settings_path = Path(directory) / "settings.json"
            with patch.object(config, "SETTINGS_PATH", settings_path):
                config.save_settings(
                    {
                        "provider_type": "custom",
                        "base_url": "https://example.invalid/v1",
                        "api_key": "sk-image-secret-value",
                        "text_api_key": "sk-text-secret-value",
                        "text_provider": "gpt",
                        "text_provider_mode": "relay",
                        "text_base_url": "https://example.invalid/v1",
                    }
                )
                stored_text = settings_path.read_text(encoding="utf-8")
                self.assertNotIn("sk-image-secret-value", stored_text)
                self.assertNotIn("sk-text-secret-value", stored_text)
                self.assertNotIn("image_size", stored_text)
                self.assertIn("api_key_encrypted", stored_text)
                public = config.get_active_settings(include_secret=False)
                self.assertIn("****", public["api_key"])
                self.assertIn("****", public["text_api_key"])
                self.assertNotIn("secret-value", public["text_api_key"])

        payload = SettingsPayload(
            provider_type="custom",
            base_url="https://example.invalid/v1",
            api_key="image-key",
            text_provider="gpt",
            text_provider_mode="relay",
            text_base_url="https://example.invalid/v1",
            text_api_key="text-key",
        )
        failed = {"ok": False, "errors": ["connection failed"]}
        with (
            patch.object(settings_routes, "_merge_existing_secret", return_value=payload.model_dump()),
            patch.object(settings_routes, "_test_image_generation", return_value=failed),
            patch.object(settings_routes, "_test_text_inference", return_value=failed),
        ):
            response = settings_routes.test_settings(payload)
        self.assertFalse(response.ok)

    def test_real_text_and_image_results_are_reported_separately(self):
        payload = SettingsPayload(
            provider_type="custom",
            base_url="https://example.invalid/v1",
            api_key="image-key",
            text_provider="gpt",
            text_provider_mode="relay",
            text_base_url="https://example.invalid/v1",
            text_api_key="text-key",
        )
        image_passed = {
            "ok": True,
            "status": "passed",
            "validation_level": "real_generation",
            "duration_seconds": 1.2,
            "errors": [],
        }
        text_passed = {
            "ok": True,
            "status": "passed",
            "validation_level": "real_inference",
            "duration_seconds": 0.4,
            "errors": [],
        }
        with (
            patch.object(settings_routes, "_merge_existing_secret", return_value=payload.model_dump()),
            patch.object(settings_routes, "_test_image_generation", return_value=image_passed),
            patch.object(settings_routes, "_test_text_inference", return_value=text_passed),
        ):
            response = settings_routes.test_settings(payload)

        self.assertTrue(response.ok)
        self.assertEqual(response.data["text"]["validation_level"], "real_inference")
        self.assertEqual(response.data["image"]["validation_level"], "real_generation")

    def test_text_and_image_probes_return_independently(self):
        payload = SettingsPayload(
            provider_type="custom",
            base_url="https://example.invalid/v1",
            api_key="image-key",
            text_provider="gpt",
            text_provider_mode="relay",
            text_base_url="https://example.invalid/v1",
            text_api_key="text-key",
        )
        with (
            patch.object(settings_routes, "_settings_for_test", return_value=(payload.model_dump(), "request")),
            patch.object(settings_routes, "validate_text_settings", return_value=[]),
            patch.object(settings_routes, "get_text_api_config", return_value={"provider": "gpt", "model": "text"}),
            patch.object(settings_routes, "_test_text_inference", return_value={"ok": True, "status": "passed"}),
        ):
            text_response = settings_routes.test_text_settings(payload)
        with (
            patch.object(settings_routes, "_settings_for_test", return_value=(payload.model_dump(), "request")),
            patch.object(settings_routes, "validate_api_settings", return_value=[]),
            patch.object(settings_routes, "_test_image_generation", return_value={"ok": False, "status": "timeout", "errors": ["timeout"]}),
        ):
            image_response = settings_routes.test_image_settings(payload)
        self.assertTrue(text_response.ok)
        self.assertFalse(image_response.ok)
        self.assertEqual(image_response.data["status"], "timeout")

    def test_saved_keys_can_be_explicitly_cleared(self):
        with tempfile.TemporaryDirectory() as directory:
            settings_path = Path(directory) / "settings.json"
            with patch.object(config, "SETTINGS_PATH", settings_path):
                config.save_settings(
                    {
                        "provider_type": "custom",
                        "base_url": "https://example.invalid/v1",
                        "api_key": "image-secret",
                        "text_provider": "gpt",
                        "text_provider_mode": "relay",
                        "text_base_url": "https://example.invalid/v1",
                        "text_api_key": "text-secret",
                    }
                )
                response = settings_routes.post_settings(
                    SettingsPayload(
                        provider_type="custom",
                        base_url="https://example.invalid/v1",
                        text_provider="gpt",
                        text_provider_mode="relay",
                        text_base_url="https://example.invalid/v1",
                        clear_saved_keys=True,
                    )
                )
                stored = settings_path.read_text(encoding="utf-8")

        self.assertTrue(response.ok)
        self.assertEqual(response.data["api_key"], "")
        self.assertEqual(response.data["text_api_key"], "")
        self.assertNotIn("api_key_encrypted", stored)
        self.assertNotIn("text_api_key_encrypted", stored)


if __name__ == "__main__":
    unittest.main()
